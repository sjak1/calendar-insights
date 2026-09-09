"""
Getting readable text out of an executive briefing document.

The EBD arrives as a blob on the event record in Oracle, or as a file on disk
in local runs, and its format is whatever the uploader happened to save. Both
the reading and the decision about what a file actually IS live here; the
generator only asks for text.

The quality gate at the bottom is what stands between the model and a document
that parsed cleanly but says nothing.
"""

import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Dict, Optional

from sqlalchemy import text

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database import engine  # noqa: E402 — shared engine from database.py
from logging_config import get_logger  # noqa: E402
from tools.extract_ebd import extract_pptx_content, format_extracted_content  # noqa: E402
from tools.agenda_config import (  # noqa: E402
    EBD_MAX_NOISE_RATIO,
    EBD_MIN_WORDS,
    MAX_DOCUMENT_CHARS,
)

logger = get_logger(__name__)

# PDF extraction (optional dependency)
try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

# DOCX extraction (optional dependency)
try:
    import docx
    from docx.table import Table as _DocxTable
    from docx.text.paragraph import Paragraph as _DocxParagraph
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
# ============================================================================
# EBD EXTRACTION FROM DATABASE
# ============================================================================

def _extract_pdf_text(pdf_path: str) -> str:
    """
    Extract text from a PDF file using pdfplumber.

    Handles arbitrary PDF layouts — extracts both free-form text and tables,
    and concatenates them page-by-page so the LLM gets a coherent view.
    """
    if not HAS_PDFPLUMBER:
        logger.warning("pdfplumber not installed. Run: pip install pdfplumber")
        return ""

    parts: list[str] = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_parts: list[str] = []

                # --- free-form text ---
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    page_parts.append(page_text.strip())

                # --- tables (if any) ---
                tables = page.extract_tables()
                for table in tables:
                    rows = []
                    for row in table:
                        cells = [
                            (cell or "").strip() for cell in row
                        ]
                        if any(cells):
                            rows.append(" | ".join(cells))
                    if rows:
                        page_parts.append("[Table]\n" + "\n".join(rows))

                if page_parts:
                    parts.append(
                        f"--- Page {page_num} ---\n" + "\n\n".join(page_parts)
                    )

        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Error extracting PDF text: {e}")
        return ""


def _extract_docx_text(docx_path: str) -> str:
    """
    Extract text from a Word document using python-docx.

    Paragraphs and tables are walked in document order rather than read from
    `.paragraphs` and `.tables` separately: those two lists lose the
    interleaving, and an EBD's tables sit under the headings that explain them.
    """
    if not HAS_DOCX:
        logger.warning("python-docx not installed. Run: pip install python-docx")
        return ""

    parts: list[str] = []
    try:
        document = docx.Document(docx_path)
        for child in document.element.body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                text = _DocxParagraph(child, document).text.strip()
                if text:
                    parts.append(text)
            elif tag == "tbl":
                rows = []
                for row in _DocxTable(child, document).rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append("[Table]\n" + "\n".join(rows))

        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        return ""


def _extract_plain_text(path: str) -> str:
    """Read a text-shaped document (txt, md, csv, json) straight off disk."""
    try:
        return Path(path).read_text(encoding="utf-8", errors="replace").strip()
    except Exception as e:
        logger.error(f"Error reading text document: {e}")
        return ""


# Formats we can turn into text. Anything outside this set is skipped with a
# warning rather than guessed at — the old code defaulted every unknown type to
# PPTX, so a Word EBD reached python-pptx and blew up.
_TEXT_SUFFIXES = frozenset({".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".rtf"})
_BINARY_SUFFIXES = frozenset({".pdf", ".pptx", ".docx"})

# Office MIME types are long and near-identical, so match on the part that
# actually distinguishes them. Order matters: first hit wins.
_CONTENT_TYPE_MARKERS = (
    ("wordprocessingml", ".docx"),
    ("presentationml", ".pptx"),
    ("pdf", ".pdf"),
    ("json", ".json"),
    ("csv", ".csv"),
    ("markdown", ".md"),
    ("text/plain", ".txt"),
)


def _sniff_document_suffix(source) -> str:
    """
    Identify a document from its own bytes. `source` is a blob or a path.

    Filenames lie: 8 of the 42 EBDs in the database carry a legacy .doc or
    .ppt extension over what is really an OOXML payload, and would otherwise
    be written off as unreadable. A magic number cannot lie the same way, so
    it is checked before the name. Returns "" when the bytes say nothing
    useful, leaving the name and MIME type to answer.
    """
    try:
        if isinstance(source, (bytes, bytearray)):
            head, archive = bytes(source[:8]), io.BytesIO(bytes(source))
        else:
            with open(source, "rb") as fh:
                head = fh.read(8)
            archive = source

        if head[:4] == b"%PDF":
            return ".pdf"
        if head[:4] != b"PK\x03\x04":
            return ""

        names = zipfile.ZipFile(archive).namelist()
        if "word/document.xml" in names:
            return ".docx"
        if any(name.startswith("ppt/slides/") for name in names):
            return ".pptx"
    except Exception as e:
        logger.debug(f"Could not sniff document type: {e}")

    return ""


def _document_suffix(filename: str = "", content_type: str = "") -> str:
    """
    Decide which extension to parse a document as, from its name or MIME type.

    The filename is trusted first — the database stores a real one, while
    content types arrive in several spellings for the same format. Returns ""
    when neither identifies something we can read.
    """
    name = (filename or "").lower()
    for suffix in _BINARY_SUFFIXES | _TEXT_SUFFIXES:
        if name.endswith(suffix):
            return suffix

    ct = (content_type or "").lower()
    for marker, suffix in _CONTENT_TYPE_MARKERS:
        if marker in ct:
            return suffix

    return ""


def _extract_document_text(path: str, suffix: str) -> tuple:
    """
    Pull text out of a document, dispatching on its suffix.

    Returns (text, extras) where extras carries whatever structural counts the
    format exposes — slides and tables for PPTX, nothing for the rest. An
    unreadable format returns empty text, which is the caller's cue to carry on
    without the document instead of failing the whole agenda.
    """
    if suffix == ".pdf":
        return _extract_pdf_text(path), {}
    if suffix == ".pptx":
        extracted = extract_pptx_content(path)
        return format_extracted_content(extracted), {
            "slide_count": extracted.get("slide_count", 0),
            "table_count": len(extracted.get("tables", [])),
        }
    if suffix == ".docx":
        return _extract_docx_text(path), {}
    if suffix in _TEXT_SUFFIXES:
        return _extract_plain_text(path), {}

    logger.warning(f"Unsupported document type '{suffix or 'unknown'}' — skipping extraction")
    return "", {}


def _fetch_ebd_from_db(event_id: str) -> Optional[Dict[str, Any]]:
    """
    Fetch EBD document from database for a given event.
    
    Args:
        event_id: The event ID to fetch EBD for
        
    Returns:
        Dict with 'raw_text' and 'has_ebd' if found, None otherwise
    """
    # VW_EVENT_DOCUMENT_REPORT.eventid holds the NUMERIC id, while callers pass
    # the UUID from x-cloud-eventid. Comparing those never matches, so every EBD
    # lookup returned "not found" regardless of whether a document was attached.
    from tools.event_resolver import resolve_numeric_event_id

    numeric_id = resolve_numeric_event_id(event_id)
    if not numeric_id:
        logger.info(f"Could not resolve a numeric id for event {event_id}; skipping EBD lookup")
        return None
    if numeric_id != event_id:
        logger.info(f"EBD lookup: {event_id} → numeric id {numeric_id}")

    try:
        with engine.connect() as conn:
            # Query for EBD document blob
            query = text("""
                SELECT document, file_name, content_type, file_size
                FROM VW_EVENT_DOCUMENT_REPORT 
                WHERE eventid = :event_id 
                AND document_category = 'Executive Briefing Document'
                AND document IS NOT NULL
                FETCH FIRST 1 ROW ONLY
            """)
            result = conn.execute(query, {"event_id": numeric_id})
            row = result.fetchone()
            
            if not row:
                logger.info(f"No EBD found in database for event: {event_id}")
                return None
            
            blob = row[0]
            filename = row[1] or "document"
            content_type = row[2] or ""
            file_size = row[3] or 0
            
            logger.info(f"Found EBD in DB: {filename} ({content_type}, {file_size} bytes)")
            
            # Determine file type and extract text
            extracted_text = ""

            sniffed = _sniff_document_suffix(blob)
            claimed = _document_suffix(filename, content_type)
            if sniffed and claimed and sniffed != claimed:
                logger.info(
                    f"EBD '{filename}' is really a {sniffed} despite its name; "
                    f"reading it as one"
                )
            suffix = sniffed or claimed
            if not suffix:
                logger.warning(
                    f"EBD '{filename}' ({content_type}) is in a format we cannot read"
                )
                return None

            # Save blob to temp file
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(blob)
                tmp_path = tmp.name

            try:
                extracted_text, _extras = _extract_document_text(tmp_path, suffix)
                if extracted_text:
                    logger.info(
                        f"Extracted {len(extracted_text)} chars from {suffix.lstrip('.').upper()}"
                    )
            finally:
                # Cleanup temp file
                Path(tmp_path).unlink(missing_ok=True)
            
            if extracted_text:
                return {
                    "raw_text": _truncate_document(extracted_text),
                    "has_ebd": True,
                    "source": "database",
                    "filename": filename,
                }

            return None

    except Exception as e:
        logger.error(f"Error fetching EBD from database: {e}", exc_info=True)
        return None


def _truncate_document(doc_text: str, max_chars: Optional[int] = None) -> str:
    """
    Truncate document text to stay within token-safe limits.

    If the text exceeds *max_chars* it is trimmed and a notice is appended
    so the LLM knows content was cut.
    """
    limit = max_chars or MAX_DOCUMENT_CHARS
    if len(doc_text) <= limit:
        return doc_text

    logger.warning(
        f"Document text truncated from {len(doc_text)} to {limit} chars"
    )
    return doc_text[:limit] + "\n\n[... document truncated due to length ...]"


def _ebd_quality_ok(extracted_text: str) -> bool:
    """Return True if extracted EBD text is usable (not garbled / too short)."""
    words = extracted_text.split()
    if len(words) < EBD_MIN_WORDS:
        logger.warning(f"EBD text too short ({len(words)} words < {EBD_MIN_WORDS}). Skipping.")
        return False
    alpha_chars = sum(1 for c in extracted_text if c.isalnum() or c.isspace())
    total_chars = len(extracted_text)
    if total_chars > 0 and (1 - alpha_chars / total_chars) > EBD_MAX_NOISE_RATIO:
        logger.warning(
            f"EBD text appears garbled (noise ratio {1 - alpha_chars / total_chars:.0%}). Skipping."
        )
        return False
    return True


